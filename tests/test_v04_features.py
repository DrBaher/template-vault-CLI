"""v0.4.0 additions: find --json, history, verify, export, --why, completion, color."""

import json
import os
import unittest
from unittest import mock

from tests._helpers import (
    CliCase, add_template, run_cli, temp_vault, tvc,
    SAMPLE_NDA_MUTUAL, SAMPLE_NDA_STARTUP_FRIENDLY,
)


class FindJsonTests(CliCase):
    def test_find_json_emits_structured_payload(self):
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL,
                         tags=["mutual"], jurisdiction=["California"],
                         summary="house mutual NDA")
            add_template(v, "nda", "yc", SAMPLE_NDA_STARTUP_FRIENDLY,
                         tags=["yc"], summary="yc-style startup NDA")
            code, out, _err = run_cli("find", "mutual", "--json")
            self.assertEqual(code, 0)
            doc = json.loads(out)
            self.assertEqual(doc["query"], "mutual")
            self.assertGreaterEqual(len(doc["results"]), 1)
            top = doc["results"][0]
            self.assertEqual(top["ref"], "nda/house")
            self.assertIn("score", top)
            self.assertIn("summary", top)
            self.assertEqual(top["tags"], ["mutual"])

    def test_find_json_empty_results(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL, summary="house")
            code, out, _err = run_cli("find", "asdfqwertyzxcvbnm", "--json")
            self.assertEqual(code, 0)
            doc = json.loads(out)
            self.assertEqual(doc["results"], [])


class HistoryTests(CliCase):
    def test_history_shows_versions_and_swaps(self):
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            add_template(v, "nda", "yc", SAMPLE_NDA_STARTUP_FRIENDLY)
            run_cli("compose", "--base", "nda/house",
                    "--as", "nda/house-startup")
            run_cli("swap", "nda/house-startup",
                    "--clause", "Term and Survival",
                    "--from", "nda/yc")
            code, out, _err = run_cli("history", "nda/house-startup")
            self.assertEqual(code, 0)
            self.assertIn("forked from nda/house", out)
            self.assertIn("v1", out)
            self.assertIn("swap", out)
            self.assertIn("Term and Survival", out)

    def test_history_json(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            code, out, _err = run_cli("history", "nda/x", "--json")
            self.assertEqual(code, 0)
            doc = json.loads(out)
            self.assertEqual(doc["ref"], "nda/x")
            self.assertIsInstance(doc["events"], list)
            self.assertTrue(any(e["kind"] == "version" for e in doc["events"]))


class VerifyTests(CliCase):
    def test_verify_records_hashes_with_update_flag(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            code, out, _err = run_cli("verify", "--update-hashes")
            self.assertEqual(code, 0)
            self.assertIn("recorded sha256", out)
            # Hash is now persisted
            meta = json.loads((v / "nda" / "x" / "meta.json").read_text())
            self.assertIn("sha256", meta["versions"][0])

    def test_verify_detects_mismatch(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            # Record current hash
            run_cli("verify", "--update-hashes")
            # Tamper with the file
            (v / "nda" / "x" / "v1.md").write_text(
                (v / "nda" / "x" / "v1.md").read_text() + "\n# TAMPERED\n"
            )
            code, out, _err = run_cli("verify")
            self.assertNotEqual(code, 0)
            self.assertIn("MISMATCH", out)

    def test_verify_strict_flags_missing_hashes(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            code, out, _err = run_cli("verify", "--strict")
            self.assertNotEqual(code, 0)
            self.assertIn("NO HASH", out)


class ExportTests(CliCase):
    def setUp(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx not installed (install with [docx] extra)")

    def test_export_writes_docx_with_h2_as_heading2(self):
        import docx as _docx
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            dest = v / "out.docx"
            code, _out, _err = run_cli(
                "export", "nda/x", "--as", "docx",
                "--output", str(dest),
            )
            self.assertEqual(code, 0)
            self.assertTrue(dest.exists())
            doc = _docx.Document(str(dest))
            styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
            # SAMPLE_NDA_MUTUAL has multiple H2 sections → at least one Heading 2.
            self.assertIn("Heading 2", styles)
            # And the H1 line ("# Mutual NDA") becomes Title.
            self.assertTrue(any(p.style.name == "Title" for p in doc.paragraphs))

    def test_export_unsupported_format_errors(self):
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL)
            code, _out, err = run_cli("export", "nda/x", "--as", "pdf")
            self.assertNotEqual(code, 0)
            self.assertIn("Unsupported export format", err)


class WhyFlagTests(CliCase):
    def test_compose_why_explains_what_happened(self):
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            code, out, _err = run_cli(
                "compose", "--base", "nda/house",
                "--as", "nda/derived", "--why",
            )
            self.assertEqual(code, 0)
            self.assertIn("[why]", out)
            self.assertIn("derived_from", out)
            self.assertIn("forked_at_parent_version", out)

    def test_swap_why_lists_clause_resolution(self):
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            add_template(v, "nda", "yc", SAMPLE_NDA_STARTUP_FRIENDLY)
            run_cli("compose", "--base", "nda/house",
                    "--as", "nda/house-startup")
            code, out, _err = run_cli(
                "swap", "nda/house-startup",
                "--clause", "Term and Survival",
                "--from", "nda/yc", "--why",
            )
            self.assertEqual(code, 0)
            self.assertIn("[why]", out)
            self.assertIn("resolved source clause", out)
            self.assertIn("clause_overrides", out)

    def test_no_why_flag_no_block(self):
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            code, out, _err = run_cli(
                "compose", "--base", "nda/house", "--as", "nda/derived",
            )
            self.assertEqual(code, 0)
            self.assertNotIn("[why]", out)


class CompletionTests(CliCase):
    def test_completion_bash(self):
        code, out, _err = run_cli("completion", "bash")
        self.assertEqual(code, 0)
        self.assertIn("_template_vault_completions", out)
        self.assertIn("complete -F", out)
        # Spot-check that subcommands the user might tab through are in the list.
        for cmd in ("init", "upload", "compose", "swap", "verify", "history", "export"):
            self.assertIn(cmd, out)

    def test_completion_zsh(self):
        code, out, _err = run_cli("completion", "zsh")
        self.assertEqual(code, 0)
        self.assertIn("compdef _template_vault", out)
        self.assertIn("_describe", out)

    def test_completion_unknown_shell_errors(self):
        # argparse choices reject unknown values before our code runs.
        code, _out, err = run_cli("completion", "fish")
        self.assertNotEqual(code, 0)
        self.assertIn("fish", err.lower())

    def test_completion_bash_includes_smart_dispatch_logic(self):
        """The bash script must contain the smart-completion branches that
        call back into `template-vault __complete`."""
        code, out, _err = run_cli("completion", "bash")
        self.assertEqual(code, 0)
        self.assertIn("template-vault __complete refs", out)
        self.assertIn("template-vault __complete versions", out)
        self.assertIn("template-vault __complete categories", out)
        # Subcommands that take a ref should be enumerated in the case stmt.
        self.assertIn("get|info|diff|history|clauses|swap", out)


class CompleteHandlerTests(CliCase):
    """The hidden __complete handler used by the shell completion scripts."""

    def test_complete_refs_lists_all_templates(self):
        from tests._helpers import (
            temp_vault, add_template,
            SAMPLE_NDA_MUTUAL, SAMPLE_LICENSING,
        )
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            add_template(v, "nda", "yc", SAMPLE_NDA_MUTUAL)
            add_template(v, "licensing", "saas", SAMPLE_LICENSING)
            code, out, _err = run_cli("__complete", "refs")
            self.assertEqual(code, 0)
            refs = out.strip().split("\n")
            self.assertIn("nda/house", refs)
            self.assertIn("nda/yc", refs)
            self.assertIn("licensing/saas", refs)

    def test_complete_categories_dedups(self):
        from tests._helpers import (
            temp_vault, add_template,
            SAMPLE_NDA_MUTUAL, SAMPLE_LICENSING,
        )
        with temp_vault() as v:
            add_template(v, "nda", "house", SAMPLE_NDA_MUTUAL)
            add_template(v, "nda", "yc", SAMPLE_NDA_MUTUAL)
            add_template(v, "licensing", "saas", SAMPLE_LICENSING)
            code, out, _err = run_cli("__complete", "categories")
            self.assertEqual(code, 0)
            cats = out.strip().split("\n")
            self.assertEqual(sorted(cats), ["licensing", "nda"])

    def test_complete_versions_lists_version_ids(self):
        from tests._helpers import (
            temp_vault, add_template,
            SAMPLE_NDA_MUTUAL,
        )
        with temp_vault() as v:
            add_template(v, "nda", "x", SAMPLE_NDA_MUTUAL, version="v1")
            # Manually add a v2 entry so the listing has more than one.
            import json as _json
            meta = _json.loads((v / "nda" / "x" / "meta.json").read_text())
            meta["versions"].append({"id": "v2", "added": "2026-05-19",
                                      "supersedes": "v1", "changelog": "x"})
            (v / "nda" / "x" / "meta.json").write_text(
                _json.dumps(meta, indent=2) + "\n"
            )
            (v / "nda" / "x" / "v2.md").write_text("# v2")
            code, out, _err = run_cli("__complete", "versions", "nda/x")
            self.assertEqual(code, 0)
            self.assertEqual(out.strip().split("\n"), ["v1", "v2"])

    def test_complete_versions_unknown_ref_silent_exit_zero(self):
        from tests._helpers import temp_vault
        with temp_vault():
            code, out, _err = run_cli("__complete", "versions", "nda/no-such")
            # Silent exit (no error during shell tab-completion).
            self.assertEqual(code, 0)
            self.assertEqual(out.strip(), "")

    def test_complete_without_vault_silent_exit_zero(self):
        """Tab-completing outside any vault should NOT print an error."""
        import tempfile, os
        with tempfile.TemporaryDirectory() as d:
            prev = os.getcwd()
            try:
                os.chdir(d)
                code, out, err = run_cli("__complete", "refs")
                self.assertEqual(code, 0)
                self.assertEqual(out, "")
                # Nothing on stderr either — silent.
                self.assertEqual(err, "")
            finally:
                os.chdir(prev)


class ColorTests(CliCase):
    def test_no_color_env_disables_ansi(self):
        # Even if we force-call _green, NO_COLOR should suppress codes.
        with mock.patch.dict(os.environ, {"NO_COLOR": "1"}, clear=False):
            self.assertEqual(tvc._green("hello"), "hello")

    def test_force_color_enables_ansi_even_off_tty(self):
        with mock.patch.dict(os.environ, {"FORCE_COLOR": "1"}, clear=False):
            os.environ.pop("NO_COLOR", None)
            self.assertIn("\033[", tvc._green("hello"))


if __name__ == "__main__":
    unittest.main()
