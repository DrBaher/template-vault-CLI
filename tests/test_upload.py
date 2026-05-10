"""upload command behavior."""

import json
import unittest
from pathlib import Path

from tests._helpers import CliCase, run_cli, temp_vault, tvc, SAMPLE_NDA_MUTUAL


class UploadTests(CliCase):
    def test_upload_writes_file_and_meta(self):
        with temp_vault() as v:
            src = v / "src.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            self.assertOk(run_cli(
                "upload", str(src),
                "--category", "nda", "--name", "house-mutual",
                "--summary", "the test NDA",
                "--non-interactive",
            ))
            t_dir = v / "nda" / "house-mutual"
            self.assertTrue((t_dir / "v1.md").exists())
            meta = json.loads((t_dir / "meta.json").read_text())
            self.assertEqual(meta["latest_version"], "v1")
            self.assertEqual(meta["category"], "nda")
            self.assertEqual(meta["versions"][0]["id"], "v1")
            self.assertEqual(meta["summary"], "the test NDA")

    def test_upload_auto_increments_version(self):
        with temp_vault() as v:
            src = v / "src.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            self.assertOk(run_cli(
                "upload", str(src), "--category", "nda", "--name", "x",
                "--summary", "s", "--non-interactive",
            ))
            self.assertOk(run_cli(
                "upload", str(src), "--category", "nda", "--name", "x",
                "--supersedes", "v1", "--summary", "s", "--non-interactive",
            ))
            meta = json.loads((v / "nda" / "x" / "meta.json").read_text())
            self.assertEqual(meta["latest_version"], "v2")
            ids = [ve["id"] for ve in meta["versions"]]
            self.assertEqual(ids, ["v1", "v2"])
            v1 = next(ve for ve in meta["versions"] if ve["id"] == "v1")
            self.assertEqual(v1.get("supersedes_by"), "v2")

    def test_upload_rejects_duplicate_version(self):
        with temp_vault() as v:
            src = v / "s.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            self.assertOk(run_cli("upload", str(src), "--category", "nda",
                                  "--name", "x", "--summary", "s",
                                  "--non-interactive"))
            code, _out, err = run_cli("upload", str(src), "--category", "nda",
                                      "--name", "x", "--version", "v1",
                                      "--summary", "s", "--non-interactive")
            self.assertNotEqual(code, 0)
            self.assertIn("already exists", err)

    def test_upload_requires_existing_file(self):
        with temp_vault():
            code, _out, err = run_cli("upload", "/no/such/file",
                                      "--category", "nda", "--name", "x",
                                      "--summary", "s", "--non-interactive")
            self.assertNotEqual(code, 0)
            self.assertIn("not found", err)

    def test_upload_records_jurisdiction_and_tags(self):
        with temp_vault() as v:
            src = v / "s.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            self.assertOk(run_cli(
                "upload", str(src), "--category", "nda", "--name", "x",
                "--summary", "s", "--non-interactive",
                "--tags", "house-style,short-form",
                "--jurisdiction", "California,Delaware",
                "--license", "MIT",
            ))
            meta = json.loads((v / "nda" / "x" / "meta.json").read_text())
            self.assertEqual(sorted(meta["jurisdiction"]), ["California", "Delaware"])
            self.assertIn("short-form", meta["tags"])
            self.assertEqual(meta["license"], "MIT")

    def test_upload_supersedes_updates_prior_pointer(self):
        with temp_vault() as v:
            src = v / "s.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            run_cli("upload", str(src), "--category", "nda", "--name", "y",
                    "--summary", "s", "--non-interactive")
            run_cli("upload", str(src), "--category", "nda", "--name", "y",
                    "--summary", "s2", "--supersedes", "v1", "--non-interactive")
            meta = json.loads((v / "nda" / "y" / "meta.json").read_text())
            v1 = next(ve for ve in meta["versions"] if ve["id"] == "v1")
            self.assertEqual(v1["supersedes_by"], "v2")
            v2 = next(ve for ve in meta["versions"] if ve["id"] == "v2")
            self.assertEqual(v2["supersedes"], "v1")

    def test_upload_outside_vault_errors(self):
        # Run from a freshly created dir without a .vault.json
        import os, tempfile
        with tempfile.TemporaryDirectory() as d:
            prev = os.getcwd()
            try:
                os.chdir(d)
                src = Path(d) / "s.md"
                src.write_text("body")
                code, _o, err = run_cli("upload", str(src),
                                        "--category", "nda", "--name", "x",
                                        "--summary", "s", "--non-interactive")
                self.assertNotEqual(code, 0)
                self.assertIn("vault", err.lower())
            finally:
                os.chdir(prev)


class UploadAmendTests(CliCase):
    def test_amend_overwrites_in_place_and_records_prior_sha(self):
        with temp_vault() as v:
            src = v / "src.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            run_cli("upload", str(src),
                    "--category", "nda", "--name", "x",
                    "--summary", "s", "--non-interactive")
            old_path = v / "nda" / "x" / "v1.md"
            old_bytes = old_path.read_bytes()
            old_sha = tvc.sha256_bytes(old_bytes)
            # Now amend with a fixed-up version of the same file.
            patched = v / "patched.md"
            patched.write_text(SAMPLE_NDA_MUTUAL.replace("two years", "two YEARS"))
            code, out, _err = run_cli(
                "upload", str(patched),
                "--category", "nda", "--name", "x",
                "--amend", "v1", "--yes-amend",
                "--non-interactive",
            )
            self.assertEqual(code, 0)
            self.assertIn("Amended", out)
            self.assertIn(old_sha, out)
            # File overwritten in place; no v2 created.
            self.assertFalse((v / "nda" / "x" / "v2.md").exists())
            self.assertIn("two YEARS", old_path.read_text())
            # meta.json's v1 changelog records the prior sha.
            meta = json.loads((v / "nda" / "x" / "meta.json").read_text())
            v1 = next(v_ for v_ in meta["versions"] if v_["id"] == "v1")
            self.assertIn("amended", v1["changelog"])
            self.assertIn(old_sha, v1["changelog"])
            self.assertIn("amended", v1)

    def test_amend_unknown_version_fails_fast(self):
        with temp_vault() as v:
            src = v / "src.md"
            src.write_text(SAMPLE_NDA_MUTUAL)
            run_cli("upload", str(src),
                    "--category", "nda", "--name", "x",
                    "--summary", "s", "--non-interactive")
            patched = v / "p.md"
            patched.write_text("x")
            code, _out, err = run_cli(
                "upload", str(patched),
                "--category", "nda", "--name", "x",
                "--amend", "v9", "--yes-amend",
                "--non-interactive",
            )
            self.assertNotEqual(code, 0)
            self.assertIn("no such version", err.lower())


class UploadDocxTests(CliCase):
    """Optional [docx] extra. If python-docx isn't installed, skip."""

    def setUp(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx not installed (install with [docx] extra)")

    def test_uploading_docx_converts_to_markdown_with_h2_headings(self):
        import docx as _docx
        with temp_vault() as v:
            src = v / "in.docx"
            doc = _docx.Document()
            doc.add_paragraph("Some preamble.")
            doc.add_paragraph("Purpose", style="Heading 2")
            doc.add_paragraph("The parties wish to evaluate.")
            doc.add_paragraph("Term and Survival", style="Heading 2")
            doc.add_paragraph("Two years from the Effective Date.")
            doc.save(str(src))
            code, _out, _err = run_cli(
                "upload", str(src),
                "--category", "nda", "--name", "from-docx",
                "--summary", "s", "--non-interactive",
            )
            self.assertEqual(code, 0)
            # Stored as .md, not .docx.
            self.assertTrue((v / "nda" / "from-docx" / "v1.md").exists())
            self.assertFalse((v / "nda" / "from-docx" / "v1.docx").exists())
            md = (v / "nda" / "from-docx" / "v1.md").read_text()
            self.assertIn("## Purpose", md)
            self.assertIn("## Term and Survival", md)
            # Clauses detect via H2 normally now:
            code, out, _err = run_cli("clauses", "nda/from-docx")
            self.assertEqual(code, 0)
            self.assertIn("Term and Survival", out)
            # Provenance recorded in changelog.
            meta = json.loads((v / "nda" / "from-docx" / "meta.json").read_text())
            self.assertIn("converted from in.docx", meta["versions"][0]["changelog"])


class UploadDocxNoExtraTests(CliCase):
    """When python-docx ISN'T installed, .docx upload errors with install hint."""

    def test_helpful_error_when_docx_extra_missing(self):
        try:
            import docx  # noqa: F401
            self.skipTest("python-docx IS installed; skipping the missing-extra path")
        except ImportError:
            pass
        with temp_vault() as v:
            src = v / "in.docx"
            src.write_bytes(b"PK\x03\x04 (fake docx, not parsed)")
            code, _out, err = run_cli(
                "upload", str(src),
                "--category", "nda", "--name", "x",
                "--summary", "s", "--non-interactive",
            )
            self.assertNotEqual(code, 0)
            self.assertIn("[docx]", err)
            self.assertIn("template-vault-cli[docx]", err)


if __name__ == "__main__":
    unittest.main()
